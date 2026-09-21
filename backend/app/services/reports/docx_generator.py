import io
import os
import urllib.request
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.schemas.domain import InspectionReportModel
from app.core.logging import logger

class DocxReportGenerator:
    @staticmethod
    def generate_docx(report_data: InspectionReportModel) -> bytes:
        doc = Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Title / Branding
        title = doc.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Legal Metrology Inspection Compliance Report")
        run.font.name = "Arial"
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 37, 55)  # Navy

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle.add_run("LM TRACE evidence based assessment record")
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGBColor(30, 64, 175)  # Royal Blue

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 1. Inspection Metadata Table
        meta_table = doc.add_table(rows=7, cols=2)
        meta_table.style = "Table Grid"
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            ("Inspection ID / Code:", f"{report_data.inspection_code} (ID: {report_data.inspection_id})"),
            ("Inspection Date & Time:", report_data.inspection_date),
            ("Inspector Details:", f"{report_data.inspector_name} (Officer ID: {report_data.officer_id})"),
            ("Inspection Location:", report_data.location),
            ("Establishment / Seller:", f"{report_data.seller_name or 'N/A'} ({report_data.business_name or 'N/A'})"),
            ("Report Version:", str(report_data.report_version)),
            ("Generated At:", report_data.generated_at),
        ]
        for idx, (lbl, val) in enumerate(meta_data):
            row = meta_table.rows[idx]
            r0 = row.cells[0].paragraphs[0].add_run(lbl)
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            r1 = row.cells[1].paragraphs[0].add_run(str(val))
            r1.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 2. Product Summary & Overall Assessment
        doc.add_heading("1 Product Information and Overall Assessment", level=2)
        p_prod = doc.add_paragraph()
        p_prod.add_run(f"Product Name: ").bold = True
        p_prod.add_run(f"{report_data.product_name} | ")
        p_prod.add_run(f"Brand: ").bold = True
        p_prod.add_run(f"{report_data.brand or 'N/A'} | ")
        p_prod.add_run(f"Category: ").bold = True
        p_prod.add_run(f"{report_data.category}\n")
        p_prod.add_run(f"Overall Compliance Status: ").bold = True
        status_run = p_prod.add_run(f"{report_data.overall_status} ")
        status_run.bold = True
        p_prod.add_run(f"| Analytical Assessment Score: ").bold = True
        p_prod.add_run(f"{report_data.score} / 100")

        finding_count = len(report_data.findings)
        failed_checks = sum(1 for item in report_data.compliance_checks if item.get("result") == "POTENTIAL_VIOLATION")
        review_checks = sum(1 for item in report_data.compliance_checks if item.get("result") in {"REVIEW", "UNVERIFIED"})
        summary_table = doc.add_table(rows=2, cols=4)
        summary_table.style = "Table Grid"
        for idx, text in enumerate(["Potential Violations", "Failed Checks", "Review or Unverified", "Evidence Items"]):
            summary_table.rows[0].cells[idx].paragraphs[0].add_run(text).bold = True
        for idx, value in enumerate([finding_count, failed_checks, review_checks, len(report_data.evidence_images)]):
            summary_table.rows[1].cells[idx].paragraphs[0].add_run(str(value))

        # 3. Declarations Matrix Table (Dual AI vs Inspector Verified Values)
        doc.add_heading("2 Statutory Declarations Verification Matrix", level=2)
        dec_headers = ["Declaration", "AI Extracted Value", "Verified Value", "Presence", "Correctness", "Status"]
        dec_table = doc.add_table(rows=1, cols=6)
        dec_table.style = "Table Grid"
        dec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = dec_table.rows[0].cells
        for i, header_text in enumerate(dec_headers):
            h_run = hdr_cells[i].paragraphs[0].add_run(header_text)
            h_run.font.bold = True
            h_run.font.size = Pt(9)

        for dec in report_data.declarations:
            row_cells = dec_table.add_row().cells
            row_cells[0].paragraphs[0].add_run(dec.get("declaration", "")).font.size = Pt(8.5)
            row_cells[1].paragraphs[0].add_run(dec.get("ai_value") or "Missing").font.size = Pt(8.5)
            row_cells[2].paragraphs[0].add_run(dec.get("verified_value") or dec.get("ai_value") or "Unverified").font.size = Pt(8.5)
            row_cells[3].paragraphs[0].add_run(dec.get("presence", "DETECTED")).font.size = Pt(8.5)
            row_cells[4].paragraphs[0].add_run(dec.get("correctness", "UNVERIFIED")).font.size = Pt(8.5)
            s_run = row_cells[5].paragraphs[0].add_run(dec.get("final_check", "UNVERIFIED"))
            s_run.font.size = Pt(8.5)
            s_run.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 4. Compliance Checks Matrix
        doc.add_heading("3 Legal Metrology Compliance Checks", level=2)
        chk_headers = ["Rule", "Check", "Field", "Input", "Expected Standard", "Result"]
        chk_table = doc.add_table(rows=1, cols=6)
        chk_table.style = "Table Grid"
        chk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(chk_headers):
            chk_table.rows[0].cells[i].paragraphs[0].add_run(h).font.bold = True

        for chk in report_data.compliance_checks:
            r = chk_table.add_row().cells
            r[0].paragraphs[0].add_run(chk.get("rule_code", "")).font.size = Pt(8.5)
            r[1].paragraphs[0].add_run(chk.get("check_type", "")).font.size = Pt(8.5)
            r[2].paragraphs[0].add_run(chk.get("field_name", "")).font.size = Pt(8.5)
            r[3].paragraphs[0].add_run(str(chk.get("input_value", ""))).font.size = Pt(8.5)
            r[4].paragraphs[0].add_run(chk.get("expected_condition", "")).font.size = Pt(8.5)
            res_run = r[5].paragraphs[0].add_run(chk.get("result", ""))
            res_run.font.size = Pt(8.5)
            res_run.font.bold = True

            explanation = chk.get("explanation")
            source = chk.get("source_reference")
            if explanation or source:
                detail = doc.add_paragraph()
                detail.paragraph_format.space_after = Pt(3)
                detail.add_run(f"{chk.get('field_name', 'Check')}: ").bold = True
                detail.add_run(str(explanation or ""))
                if source:
                    detail.add_run(f" Source: {source}").italic = True

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 4. Violation summary
        doc.add_heading("4 Violation and Review Summary", level=2)
        if report_data.findings:
            finding_table = doc.add_table(rows=1, cols=5)
            finding_table.style = "Table Grid"
            for idx, heading in enumerate(["Type", "Severity", "Status", "Description", "Inspector Comment"]):
                finding_table.rows[0].cells[idx].paragraphs[0].add_run(heading).bold = True
            for finding in report_data.findings:
                cells = finding_table.add_row().cells
                values = [
                    finding.get("type", "FINDING"), finding.get("severity", "MEDIUM"),
                    finding.get("status", "AI_DETECTED"), finding.get("ai_explanation") or finding.get("explanation") or "",
                    finding.get("inspector_comment") or "Pending inspector disposition",
                ]
                for idx, value in enumerate(values):
                    cells[idx].paragraphs[0].add_run(str(value)).font.size = Pt(8)
        else:
            doc.add_paragraph("No potential violations were recorded. Confirm that no checks remain under review or unverified before enforcement closure.")

        # 5. Visual Evidence and forensic crops
        if report_data.evidence_images:
            doc.add_heading("5 Visual Evidence and Forensic Crops", level=2)
            ev_count = 0
            for ev in report_data.evidence_images:
                if ev_count >= 6:
                    break
                sec_url = ev.get("cloudinary_secure_url")
                local_path = ev.get("crop_path")
                ev_type = ev.get("evidence_type", "EVIDENCE")
                sha = ev.get("sha256", "N/A")
                desc = ev.get("description", "")

                img_bytes = None
                if sec_url:
                    try:
                        req = urllib.request.Request(sec_url, headers={"User-Agent": "LM-Trace-ReportGen/1.0"})
                        with urllib.request.urlopen(req, timeout=3) as resp:
                            img_bytes = resp.read()
                    except Exception as dl_err:
                        logger.debug("Could not fetch remote Cloudinary image %s: %s", sec_url, dl_err)

                if not img_bytes and local_path and os.path.exists(local_path):
                    try:
                        with open(local_path, "rb") as f:
                            img_bytes = f.read()
                    except Exception:
                        pass

                if img_bytes:
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run_img = p_img.add_run()
                        run_img.add_picture(io.BytesIO(img_bytes), width=Inches(3.5))

                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap_run = p_cap.add_run(f"[{ev_type}] {desc} | SHA-256: {sha[:16]}...")
                        cap_run.font.size = Pt(8)
                        cap_run.font.italic = True
                        ev_count += 1
                    except Exception as embed_err:
                        logger.warning("Could not embed evidence image in DOCX: %s", embed_err)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 6. Inspector Remarks & Disclaimers
        doc.add_heading("6 Inspector Remarks and Regulatory Disclaimer", level=2)
        p_rem = doc.add_paragraph()
        p_rem.add_run("Inspector Observations & Orders:\n").bold = True
        p_rem.add_run(report_data.inspector_remarks or "Inspection completed with visual evidence and digital audit trail recorded.")

        p_disc = doc.add_paragraph()
        p_disc.paragraph_format.space_before = Pt(12)
        d_run = p_disc.add_run(f"LEGAL DISCLAIMER: {report_data.disclaimer}")
        d_run.font.size = Pt(8)
        d_run.font.italic = True
        d_run.font.color.rgb = RGBColor(100, 116, 139)

        # Signature Block
        sig_p = doc.add_paragraph()
        sig_p.paragraph_format.space_before = Pt(20)
        sig_p.add_run("____________________________\n").bold = True
        sig_p.add_run(f"Verified & Authorized by: {report_data.inspector_name}\n")
        sig_p.add_run(f"Officer ID: {report_data.officer_id} | Date: {report_data.inspection_date}\n")
        sig_p.add_run("Legal Metrology Department, Government of India")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

docx_report_generator = DocxReportGenerator()
